from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Inches
import os


# =========================
# LOAD SOURCE DOCX
# =========================
src_doc = Document("C:/Users/Admin/Desktop/Toán học 2026 - Đề số 0.docx")

# =========================
# CREATE OUTPUT DOCX
# =========================
out_doc = Document()


# =========================
# ITERATE BLOCKS
# =========================
def iter_block_items(parent):

    parent_elm = parent.element.body

    for child in parent_elm.iterchildren():

        if child.tag.endswith('}p'):
            yield Paragraph(child, parent)

        elif child.tag.endswith('}tbl'):
            yield Table(child, parent)


# =========================
# IMAGE DETECTION
# =========================
def get_images_from_paragraph(paragraph):

    image_ids = []

    for run in paragraph.runs:

        drawing_elements = run._element.xpath('.//w:drawing')

        if drawing_elements:

            blips = run._element.xpath('.//a:blip')

            for blip in blips:

                r_embed = blip.get(
                    '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
                )

                if r_embed:
                    image_ids.append(r_embed)

    return image_ids


# =========================
# COPY CONTENT
# =========================
for block in iter_block_items(src_doc):

    # ---------------------
    # PARAGRAPH
    # ---------------------
    if isinstance(block, Paragraph):

        text = block.text.strip()

        if text:
            out_doc.add_paragraph(text)

        # -----------------
        # COPY IMAGES
        # -----------------
        image_ids = get_images_from_paragraph(block)

        for image_id in image_ids:

            image_part = src_doc.part.related_parts[image_id]

            image_bytes = image_part.blob

            image_name = os.path.basename(str(image_part.partname))

            temp_path = f"temp_{image_name}"

            # Save temporary image
            with open(temp_path, "wb") as f:
                f.write(image_bytes)

            # Insert image into new doc
            out_doc.add_picture(temp_path, width=Inches(4))

            # Remove temp file
            os.remove(temp_path)

    # ---------------------
    # TABLE
    # ---------------------
    elif isinstance(block, Table):

        rows = len(block.rows)
        cols = len(block.columns)

        new_table = out_doc.add_table(rows=rows, cols=cols)

        for i, row in enumerate(block.rows):

            for j, cell in enumerate(row.cells):

                new_table.cell(i, j).text = cell.text


# =========================
# SAVE OUTPUT
# =========================
out_doc.save("output.docx")

print("DONE")